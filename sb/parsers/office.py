"""Read-only DOCX/XLSX parsers. Formulas are cached values or unavailable, never evaluated."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from typing import Final

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

from sb.parsers import FormatUnsupported, InvalidInput, LimitExceeded

DOCX_MEDIA_TYPE: Final = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MEDIA_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMULA_CACHE_TOKEN: Final = "formula-cache"
FORMULA_UNAVAILABLE_TOKEN: Final = "formula-unavailable"
_ENCRYPTED_MEMBERS: Final = frozenset({"encryptioninfo", "encryptedpackage"})
_OPEN_ERRORS: Final = (
    OSError,
    KeyError,
    ValueError,
    zipfile.BadZipFile,
    PackageNotFoundError,
    InvalidFileException,
)


@dataclass(frozen=True, slots=True)
class OfficeParseLimits:
    """Admission bounds for one OOXML package."""

    zip_members: int = 10_000
    zip_member_bytes: int = 67_108_864
    zip_uncompressed_bytes: int = 536_870_912
    zip_ratio: int = 100
    worksheets: int = 100
    rows: int = 1_000_000
    cells: int = 10_000_000


DEFAULT_OFFICE_PARSE_LIMITS: Final = OfficeParseLimits()


def _member_violation(info: zipfile.ZipInfo, limits: OfficeParseLimits, uncompressed: int) -> None:
    name = info.filename.replace("\\", "/")
    lowered = name.lower()
    if name.startswith("/") or any(part == ".." for part in name.split("/")):
        raise InvalidInput("office zip member path")
    if info.file_size > limits.zip_member_bytes or uncompressed > limits.zip_uncompressed_bytes:
        raise LimitExceeded("office zip size")
    if info.compress_size > 0 and info.file_size / info.compress_size > limits.zip_ratio:
        raise LimitExceeded("office zip compression ratio")
    if lowered.rsplit("/", 1)[-1] in _ENCRYPTED_MEMBERS:
        raise InvalidInput("office encrypted package")
    if lowered.endswith("vbaproject.bin") or lowered.startswith("xl/externallinks/"):
        raise FormatUnsupported("office macros or external links")


def admit_office_package(payload: bytes, limits: OfficeParseLimits, required_member: str) -> None:
    """Admit one OOXML package against ZIP, encryption, macro, and link limits."""
    if not payload:
        raise InvalidInput("empty office payload")
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as exc:
        raise InvalidInput("corrupt office zip") from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > limits.zip_members:
            raise LimitExceeded("office zip member count")
        uncompressed = 0
        rels: list[str] = []
        found_required = False
        for info in infos:
            uncompressed += info.file_size
            _member_violation(info, limits, uncompressed)
            name = info.filename.replace("\\", "/")
            lowered = name.lower()
            found_required = found_required or lowered == required_member
            if lowered.endswith(".rels"):
                rels.append(name)
        if not found_required:
            raise InvalidInput("office missing required member")
        for name in rels:
            body = archive.read(name)
            if b'TargetMode="External"' in body or b"TargetMode='External'" in body:
                raise FormatUnsupported("office external relationship")


def _table_text(table: Table) -> str:
    rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return "\n".join(row for row in rows if row.strip())


def _open_xlsx(payload: bytes, *, data_only: bool) -> OpenpyxlWorkbook:
    return load_workbook(
        io.BytesIO(payload),
        read_only=True,
        data_only=data_only,
        keep_vba=False,
        keep_links=False,
    )


@dataclass(frozen=True, slots=True)
class DocxDocumentParser:
    """Extract DOCX paragraphs and tables as visible text."""

    limits: OfficeParseLimits = DEFAULT_OFFICE_PARSE_LIMITS

    def parse(self, payload: bytes) -> str:
        admit_office_package(payload, self.limits, "word/document.xml")
        try:
            document = Document(io.BytesIO(payload))
        except _OPEN_ERRORS as exc:
            raise InvalidInput("corrupt docx") from exc
        texts: list[str] = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
            else:
                text = _table_text(block)
            if text:
                texts.append(text)
        return "\n".join(texts)


@dataclass(frozen=True, slots=True)
class XlsxDocumentParser:
    """Extract XLSX sheets and cells. Formulas stay cached or unavailable."""

    limits: OfficeParseLimits = DEFAULT_OFFICE_PARSE_LIMITS

    def parse(self, payload: bytes) -> str:
        admit_office_package(payload, self.limits, "xl/workbook.xml")
        try:
            return _extract_xlsx(payload, self.limits)
        except _OPEN_ERRORS as exc:
            raise InvalidInput("corrupt xlsx") from exc


def _formula_text(key: str, cache: dict[str, str]) -> str:
    cached = cache.get(key)
    return (
        FORMULA_UNAVAILABLE_TOKEN
        if cached is None or cached == ""
        else f"{FORMULA_CACHE_TOKEN}\t{cached}"
    )


def _cached_formulas(payload: bytes, keys: list[str]) -> dict[str, str]:
    if not keys:
        return {}
    wanted = set(keys)
    workbook = _open_xlsx(payload, data_only=True)
    found: dict[str, str] = {}
    try:
        for sheet in workbook.worksheets:
            title = sheet.title or "Sheet"
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value in {None, ""} or not hasattr(cell, "coordinate"):
                        continue
                    key = f"{title}!{cell.coordinate}"
                    if key in wanted:
                        found[key] = str(cell.value)
    finally:
        workbook.close()
    return found


def _extract_xlsx(payload: bytes, limits: OfficeParseLimits) -> str:
    workbook = _open_xlsx(payload, data_only=False)
    try:
        if len(workbook.worksheets) > limits.worksheets:
            raise LimitExceeded("xlsx worksheet count")
        pending: list[tuple[str, str | None]] = []
        formula_keys: list[str] = []
        rows = 0
        cells = 0
        for sheet in workbook.worksheets:
            title = sheet.title or "Sheet"
            pending.append((title, title))
            for row in sheet.iter_rows():
                rows += 1
                if rows > limits.rows:
                    raise LimitExceeded("xlsx row count")
                for cell in row:
                    if cell.value is None:
                        continue
                    cells += 1
                    if cells > limits.cells:
                        raise LimitExceeded("xlsx cell count")
                    key = f"{title}!{cell.coordinate}"
                    if cell.data_type == "f":
                        pending.append((key, None))
                        formula_keys.append(key)
                    else:
                        pending.append((key, str(cell.value)))
    finally:
        workbook.close()
    cache = _cached_formulas(payload, formula_keys)
    return "\n".join(_formula_text(key, cache) if text is None else text for key, text in pending)


__all__ = (
    "DOCX_MEDIA_TYPE",
    "XLSX_MEDIA_TYPE",
    "DocxDocumentParser",
    "OfficeParseLimits",
    "XlsxDocumentParser",
)

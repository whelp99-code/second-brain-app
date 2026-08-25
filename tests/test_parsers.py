"""Given/When/Then tests for the suffix-routed document parsers."""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable
from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook

from sb.parsers import (
    InvalidInput,
    UnsupportedSuffix,
    parser_for,
)
from sb.parsers.hwpx import HwpxDocumentParser


def _text_pdf(*pages: str) -> bytes:
    """Synthesize a minimal extractable PDF at test runtime."""

    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    page_ids = [4 + index * 2 for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{kids}] /Count {len(pages)} >> endobj\n".encode(),
        b"3 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    for page_id, text in zip(page_ids, pages, strict=True):
        content_id = page_id + 1
        stream = f"BT /F1 12 Tf 72 120 Td ({escape(text)}) Tj ET".encode("ascii")
        objects.append(
            (
                f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
                f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R >> >> >> endobj\n"
            ).encode()
        )
        objects.append(
            f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode()
            + stream
            + b"\nendstream\nendobj\n"
        )
    header = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    body = bytearray(header)
    offsets: list[int] = []
    for obj in objects:
        offsets.append(len(body))
        body.extend(obj)
        if not obj.endswith(b"\n"):
            body.extend(b"\n")
    xref_pos = len(body)
    xref = [f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()]
    xref.extend(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(body) + b"".join(xref) + trailer


def _hwpx(*texts: str) -> bytes:
    """Synthesize a bounded two-section-capable HWPX ZIP at test runtime."""
    section = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section"'
        ' xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
        "<hp:p><hp:run><hp:t>{text}</hp:t></hp:run></hp:p></hs:sec>"
    )
    members = {
        "mimetype": b"application/hwp+zip",
        "Contents/content.hpf": b"<hpf/>",
        **{
            f"Contents/section{index}.xml": section.format(text=text).encode("utf-8")
            for index, text in enumerate(texts)
        },
    }
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def _docx_payload() -> bytes:
    document = Document()
    _ = document.add_paragraph("Office visible body")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_payload() -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Metrics"
    sheet["A1"] = "hello-sheet"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _payload_md() -> bytes:
    return b"# heading\nvisible markdown\n"


def _payload_txt() -> bytes:
    return b"plain visible text\n"


def _payload_json() -> bytes:
    return b'{"z": 1, "a": "visible-json"}\n'


def _payload_yaml() -> bytes:
    return b"z: 1\na: visible-yaml\n"


def _payload_pdf() -> bytes:
    return _text_pdf("visible-pdf-page")


def _payload_hwpx() -> bytes:
    return _hwpx("visible-hwpx-section")


SUPPORTED: tuple[tuple[str, Callable[[], bytes]], ...] = (
    (".pdf", _payload_pdf),
    (".hwpx", _payload_hwpx),
    (".docx", _docx_payload),
    (".xlsx", _xlsx_payload),
    (".md", _payload_md),
    (".txt", _payload_txt),
    (".json", _payload_json),
    (".yaml", _payload_yaml),
)


@pytest.mark.parametrize(("suffix", "factory"), SUPPORTED)
def test_parse_returns_nonempty_text_when_suffix_is_supported(
    suffix: str, factory: Callable[[], bytes]
) -> None:
    # Given: a runtime-synthesized fixture for one registered suffix.
    payload = factory()

    # When: the suffix registry routes the payload.
    text = parser_for(suffix)(payload)

    # Then: the parser yields non-empty visible text.
    assert isinstance(text, str)
    assert text.strip() != ""


def test_parser_for_raises_unsupported_suffix_when_unknown() -> None:
    # Given: a suffix that is not in the closed registry.
    # When: the registry is asked for a parser.
    # Then: it raises the typed UnsupportedSuffix error.
    with pytest.raises(UnsupportedSuffix) as exc_info:
        _ = parser_for(".bin")
    assert exc_info.value.suffix == ".bin"


def test_hwpx_parser_raises_invalid_input_when_zip_is_truncated() -> None:
    # Given: a truncated ZIP presented as HWPX.
    parser = HwpxDocumentParser()

    # When: the constrained parser reads the payload.
    # Then: it raises the documented InvalidInput error.
    with pytest.raises(InvalidInput):
        _ = parser.parse(b"PK\x03\x04truncated")


def test_registry_raises_invalid_input_when_hwpx_is_corrupt() -> None:
    # Given: a registry-routed HWPX parser and a corrupt section package.
    parse = parser_for(".hwpx")
    section = b"<hs:sec><hp:p>"
    members = {
        "mimetype": b"application/hwp+zip",
        "Contents/section0.xml": section,
    }
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    payload = buffer.getvalue()

    # When: the registry surfaces the parser failure.
    # Then: the typed InvalidInput error is not swallowed.
    with pytest.raises(InvalidInput):
        _ = parse(payload)


def test_json_parser_emits_canonical_object_when_keys_are_unsorted() -> None:
    # Given: a JSON object whose keys are not sorted.
    payload = b'{"z":1,"a":2}'

    # When: the structured parser loads the object.
    text = parser_for(".json")(payload)

    # Then: the machine-consumed canonical form is sorted and compact.
    assert json.loads(text) == {"a": 2, "z": 1}
    assert text == '{"a":2,"z":1}'


_NON_STRING_KEY_REASON = "structured mapping key is not a string"


def test_yaml_parser_raises_invalid_input_when_key_is_numeric() -> None:
    # Given: a YAML mapping whose key is an integer scalar.
    payload = b"1: hello\n"

    # When: the structured parser loads the document.
    # Then: the non-string key is rejected with the stable typed reason.
    with pytest.raises(InvalidInput) as exc_info:
        _ = parser_for(".yaml")(payload)
    assert exc_info.value.reason == _NON_STRING_KEY_REASON


def test_yaml_parser_raises_invalid_input_when_key_is_boolean() -> None:
    # Given: a YAML mapping whose key is a boolean scalar.
    payload = b"true: yes\n"

    # When: the structured parser loads the document.
    # Then: the non-string key is rejected with the stable typed reason.
    with pytest.raises(InvalidInput) as exc_info:
        _ = parser_for(".yaml")(payload)
    assert exc_info.value.reason == _NON_STRING_KEY_REASON
